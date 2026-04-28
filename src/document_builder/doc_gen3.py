from __future__ import annotations
from pathlib import Path
from warnings import deprecated
import xml.etree.ElementTree as ET

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

@deprecated(reason="This is a demo function for chapter 3 generation. It may be removed or refactored in the future.")
def _shade_cell(cell, fill_hex: str) -> None:
    tc_pr = cell._element.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill_hex)
    tc_pr.append(shd)

@deprecated(reason="This is a demo function for chapter 3 generation. It may be removed or refactored in the future.")

def _set_cell_text(cell, text: str, *, bold: bool = False, center: bool = False) -> None:
    cell.text = ""
    p = cell.paragraphs[0]
    if center:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(text if text is not None else "")
    run.bold = bold

@deprecated(reason="This is a demo function for chapter 3 generation. It may be removed or refactored in the future.")

def render_ch3_from_xml(xml_path: str | Path, out_docx: str | Path) -> None:
    xml_path = Path(xml_path)
    out_docx = Path(out_docx)

    root = ET.parse(str(xml_path)).getroot()
    ch3 = root.find(".//paragraph[@number='3']")
    if ch3 is None:
        raise ValueError("Chapter 3 paragraph (number='3') not found in XML")

    title = (ch3.get("title") or "TASK VERSION TABLE").strip()
    table_node = ch3.find("./table")
    if table_node is None:
        raise ValueError("Chapter 3 <table> not found in XML")

    # Columns (optional) — fallback to default
    cols_attr = (table_node.get("columns") or "").strip()
    columns = [c.strip() for c in cols_attr.split(",") if c.strip()] or ["Name", "Type", "Version", "Modified"]

    doc = Document()
    doc.add_heading(f"3. {title}", level=1)
    doc.add_paragraph("")

    table = doc.add_table(rows=1, cols=len(columns))
    table.style = "Table Grid"

    # Header
    hdr = table.rows[0].cells
    for i, col in enumerate(columns):
        _set_cell_text(hdr[i], col, bold=True, center=True)
        _shade_cell(hdr[i], "D9D9D9")

    # Data rows (attributes)
    for row in table_node.findall("./row"):
        name = (row.get("name") or "").strip()
        typ = (row.get("type") or "").strip()
        ver = (row.get("version") or "").strip()
        mod = (row.get("modified") or "N/A").strip()

        r = table.add_row().cells
        _set_cell_text(r[0], name, center=False)
        _set_cell_text(r[1], typ, center=True)
        _set_cell_text(r[2], ver, center=True)
        _set_cell_text(r[3], mod, center=True)

    out_docx.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(out_docx))
