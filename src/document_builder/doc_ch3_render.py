from warnings import deprecated

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
import xml.etree.ElementTree as ET
from pathlib import Path

@deprecated(reason="This is a demo function for chapter 3 generation. It may be removed or refactored in the future.")

def render_ch3_from_xml(xml_path: Path, out_docx: Path):
    root = ET.parse(str(xml_path)).getroot()
    ch3 = root.find(".//paragraph[@number='3']")

    if ch3 is None:
        raise ValueError("Chapter 3 not found in XML")

    table_node = ch3.find("./table")
    if table_node is None:
        raise ValueError("Chapter 3 table not found in XML")

    def _shade_cell(cell, fill_hex: str) -> None:
        tc_pr = cell._element.get_or_add_tcPr()
        shd = OxmlElement("w:shd")
        shd.set(qn("w:val"), "clear")
        shd.set(qn("w:color"), "auto")
        shd.set(qn("w:fill"), fill_hex)
        tc_pr.append(shd)

    def _set_cell_text(cell, text: str, *, bold: bool = False, center: bool = False) -> None:
        cell.text = ""
        paragraph = cell.paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER if center else WD_ALIGN_PARAGRAPH.LEFT
        run = paragraph.add_run(text)
        run.bold = bold

    doc = Document()

    # Titolo capitolo
    doc.add_heading(
        f"3 {ch3.get('title')}",
        level=1
    )

    table = doc.add_table(rows=1, cols=4)
    table.style = "Table Grid"

    headers = ["Name", "Type", "Version", "Modified"]
    for i, h in enumerate(headers):
        cell = table.rows[0].cells[i]
        _set_cell_text(cell, h, bold=True, center=True)
        _shade_cell(cell, "D9D9D9")

    rows = table_node.findall("./task") or table_node.findall("./row")
    for row in rows:
        cells = table.add_row().cells
        _set_cell_text(cells[0], row.get("name", ""), center=False)
        _set_cell_text(cells[1], row.get("type", ""), center=True)
        _set_cell_text(cells[2], row.get("version", ""), center=True)
        _set_cell_text(cells[3], row.get("modified", "N/A"), center=True)

    out_docx.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(out_docx))
