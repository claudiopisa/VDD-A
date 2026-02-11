from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
import xml.etree.ElementTree as ET
from pathlib import Path

def render_ch3_from_xml(xml_path: Path, out_docx: Path):
    root = ET.parse(xml_path).getroot()
    ch3 = root.find(".//paragraph[@number='3']")

    if ch3 is None:
        raise ValueError("Chapter 3 not found in XML")

    doc = Document()

    # Titolo capitolo
    doc.add_heading(
        f"3 {ch3.get('title')}",
        level=1
    )

    table = doc.add_table(rows=1, cols=4)
    table.style = "Table Grid"

    headers = ["Name", "Type", "Version", "Modified"]
    for i, h in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.text = h
        cell.paragraphs[0].runs[0].bold = True
        cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER

    for row in ch3.findall("./table/row"):
        cells = table.add_row().cells
        cells[0].text = row.get("name", "")
        cells[1].text = row.get("type", "")
        cells[2].text = row.get("version", "")
        cells[3].text = row.get("modified", "")

        for c in cells:
            c.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER

    out_docx.parent.mkdir(parents=True, exist_ok=True)
    doc.save(out_docx)
