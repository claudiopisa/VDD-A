from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

# Create document
doc = Document()

# Main title
title = doc.add_heading('2  LIST OF WSPHS+ SW FILES AND RELEVANT VERSIONS', level=1)

# Subtitle
subtitle = doc.add_heading('2.1 NSPC_1.9.0.0\\SWPG\\INC_PROT', level=2)

# Blank line
doc.add_paragraph()

# Data
data = [
    ('APLDiag.h', '1.3'),
    ('bbfatalmsg.h', '1.1'),
    ('Cdb.h', '1.9.0'),
    ('CDBandBufferManager.h', '1.05'),
    ('CdbRedundancy.h', '1.01'),
    ('CmDef.h', '1.1'),
    ('CmMngmt.h', '1.9'),
    ('CodedSaiTypes.h', '1.02'),
    ('CommType.h', '4.01'),
    ('Connection.h', '1.17'),
    ('DiaComSS.h', '3.5'),
    ('EaiqUtils.h', '1.1'),
    ('EaoqUtils.h', '1.1'),
    ('ETCSIdType.h', '1.0.0'),
]

# Create table
table = doc.add_table(rows=1, cols=2)
table.style = 'Table Grid'

# Header
header_cells = table.rows[0].cells
header_cells[0].text = 'File Name'
header_cells[1].text = 'File Version'

# Format header (bold, centered, gray background)
for cell in header_cells:
    cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    cell.paragraphs[0].runs[0].font.bold = True
    
    # Gray background
    shading = OxmlElement('w:shd')
    shading.set(qn('w:fill'), 'D9D9D9')
    cell._element.get_or_add_tcPr().append(shading)

# Add data rows
for file_name, version in data:
    row = table.add_row()
    row.cells[0].text = file_name
    row.cells[1].text = version
    row.cells[1].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER

# Save
doc.save('file_versions2.docx')
print("Document created!")