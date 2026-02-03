from pathlib import Path
from xml.etree import ElementTree as ET
from docx import Document


def render_ch2(xml_path: str | Path, out_docx: str | Path):
    xml_path = Path(xml_path)
    out_docx = Path(out_docx)

    # read XML root
    root = ET.parse(str(xml_path)).getroot()

    # trova paragrafo indicato number="2"
    ch2 = root.find(".//paragraph[@number='2']")
    if ch2 is None:
        raise ValueError("Chapter 2 paragraph (number='2') not found in XML")

    doc = Document()

     # Titolo capitolo
    ch2_number = ch2.get("number", "2")
    ch2_title = ch2.get("title", "")
    doc.add_heading(f"{ch2_number}. {ch2_title}".strip(), level=1)

    # per ogni sottoparagrafo
    for sp in ch2.findall("./subparagraph"):
        sp_number = sp.get("number", "")
        sp_title = sp.get("title", "")
        doc.add_heading(f"{sp_number} {sp_title}".strip(), level=2) 

        
        # Ogni sottopar  può avere una o più table
        for t in sp.findall("./table"):
            # leggi colonne
            columns = []
            for c in t.findall("./columns/column"):
                key = c.get("key", "")
                label = (c.text or "").strip()
                columns.append((key, label))

            if not columns:
                # se non ci sono colonne, salta
                continue

            # crea tabella Word (prima riga sarebbe header)
            wtable = doc.add_table(rows=1, cols=len(columns))
            hdr_cells = wtable.rows[0].cells
            for j, (_, label) in enumerate(columns):
                hdr_cells[j].text = label

            # righe
            for r in t.findall("./rows/row"):
                row_cells = wtable.add_row().cells
                # mappa cell key in value
                cell_map = {}

                for cell in r.findall("./cell"):
                    k = cell.get("key", "")
                    v = (cell.text or "").strip()
                    cell_map[k] = v

                 # scrivi le celle seguendo l’ordine delle columns
                for j, (key, _) in enumerate(columns):
                    row_cells[j].text = cell_map.get(key, "")

            

            doc.add_paragraph("")  # spazio dopo la tabella

    out_docx.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(out_docx))


def render_ch2_new(xml_path: str | Path, out_docx: str | Path):
    xml_path = Path(xml_path)
    out_docx = Path(out_docx)
    
    # read XML root
    root = ET.parse(str(xml_path)).getroot()

    # trova paragrafo indicato number="2"
    ch2 = root.find(".//paragraph[@number='2']")
    if ch2 is None:
        raise ValueError("Chapter 2 paragraph (number='2') not found in XML")

    doc = Document()

     # Titolo capitolo
    ch2_number = ch2.get("number", "2")
    ch2_title = ch2.get("title", "")
    doc.add_heading(f"{ch2_number}. {ch2_title}".strip(), level=1)

    # per ogni sottoparagrafo
    for sp in ch2.findall("./subparagraph"):
        sp_number = sp.get("number", "")
        sp_title = sp.get("title", "")
        doc.add_heading(f"{sp_number} {sp_title}".strip(), level=2) #scrivi titolo sottoparagrafo

    
        # crea tabella Word (prima riga sarebbe header)
        wtable = doc.add_table(rows=1, cols=2)
        hdr_cells = wtable.rows[0].cells
        hdr_cells[0].text = "File name"
        hdr_cells[1].text = "File version"

        """for r in sp.findall("./table/row"):
            row_cells = wtable.add_row().cells
            # mappa cell key in value
            cell_map = {} 

            name = r.find("./name")
            version = r.find("./version")
            print("DEBUG | ", cell_map)

            cell_map[name.text] = version.text.strip()

            print("DEBUG | ", cell_map)

            for i, key in enumerate(cell_map.keys()):
                row_cells[i].text = cell_map.get(key)

            #row_cells[0].text = cell_map.get("name", "")
            #row_cells[1].text = cell_map.get("version", "")

            doc.add_paragraph("")"""
        
        for r in sp.findall("./table/row"):
            name_el = r.find("./name")
            ver_el = r.find("./version")

            file_name = (name_el.text or "").strip() if name_el is not None else ""
            file_ver  = (ver_el.text or "").strip() if ver_el is not None else ""

            row_cells = wtable.add_row().cells
            row_cells[0].text = file_name
            row_cells[1].text = file_ver

        doc.add_paragraph("")  # UNA riga vuota tra una tabella e la successiva

    out_docx.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(out_docx))



#main

xml_path = Path("chapter2_new.xml")
out_docx = Path("chapter2_new.docx")
#render_ch2(xml_path, out_docx)  
render_ch2_new(xml_path, out_docx)