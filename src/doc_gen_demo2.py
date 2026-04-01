from dataclasses import dataclass
from pathlib import Path
import configparser

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


# -------------------------
# MODEL
# -------------------------
@dataclass
class TaskRow:
    name: str
    type: str
    version: str
    modified: str = "N.A."  # per ora non hai baseline precedente


# -------------------------
# PARSER IMGCONF (kernel interno)
# -------------------------
def load_imgconf(imgconf_path: str | Path) -> configparser.ConfigParser:
    imgconf_path = Path(imgconf_path)

    parser = configparser.ConfigParser(
        interpolation=None,
        comment_prefixes=(";", "#", "//"),
        inline_comment_prefixes=(";", "#", "//"),
        delimiters=("=",),
        strict=False,
    )
    parser.optionxform = str  # preserva case
    with imgconf_path.open("r", encoding="utf-8", errors="ignore") as f:
        parser.read_file(f)
    return parser


def read_tasks_kernel_internal(imgconf_path: str | Path, nspc_root: str | Path) -> list[TaskRow]:
    imgconf_path = Path(imgconf_path)
    nspc_root = Path(nspc_root)

    ini = load_imgconf(imgconf_path)

    if "Settings" not in ini:
        raise ValueError("Missing [Settings] section in Imgconf.ini")

    s = ini["Settings"]
    tasks: list[TaskRow] = []

    # --- System tasks (BOOT/BOOTAP/Loader/Kernel) ---
    # Boot
    if "FileBootVer" in s:
        tasks.append(TaskRow(name="BOOT", type="SYSTEM", version=s.get("FileBootVer", "").strip()))

    # BootAP
    if "FileBootAPVer" in s:
        tasks.append(TaskRow(name="BOOTAP", type="SYSTEM", version=s.get("FileBootAPVer", "").strip()))

    # Loader: il legacy prende un filename da FileLoaderVer e poi legge la versione dal contenuto del file :contentReference[oaicite:4]{index=4}
    loader_ver_file = s.get("FileLoaderVer", "").strip()
    if loader_ver_file:
        loader_file = nspc_root / loader_ver_file
        # TODO: sostituisci con la tua extract_version_smart()
        loader_version = "<TODO extract from file>" if loader_file.exists() else "<MISSING FILE>"
        tasks.append(TaskRow(name="Loader", type="SYSTEM", version=loader_version))

    # Kernel: RelKernel :contentReference[oaicite:5]{index=5}
    if "RelKernel" in s:
        tasks.append(TaskRow(name="Kernel", type="SYSTEM", version=s.get("RelKernel", "").strip()))

    # --- BSP/app tasks (NumTask, TipoTask{i}, RelTask{i}, V1_FileTask{i}...) :contentReference[oaicite:6]{index=6}
    n_task = int(s.get("NumTask", "0"))
    for i in range(1, n_task + 1):
        ttype = s.get(f"TipoTask{i}", "").strip()
        if not ttype:
            continue
        if ttype.upper() in ("NO_SCHED", "RBC"):  # legacy skip :contentReference[oaicite:7]{index=7}
            continue

        version = s.get(f"RelTask{i}", "").strip()
        v1 = s.get(f"V1_FileTask{i}", "").strip()

        # legacy: name = nome eseguibile senza estensione preso da V1 path :contentReference[oaicite:8]{index=8}
        name = Path(v1).name
        if "." in name:
            name = name.split(".", 1)[0]

        tasks.append(TaskRow(name=name, type=ttype, version=version))

    return tasks


# -------------------------
# DOCX TABLE (stile come il tuo capitolo 2)
# -------------------------
def shade_cell(cell, fill_hex: str = "D9D9D9"):
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), fill_hex)
    cell._element.get_or_add_tcPr().append(shading)


def build_ch3_docx(tasks: list[TaskRow], out_docx: str | Path):
    out_docx = Path(out_docx)

    doc = Document()
    doc.add_heading("3. TASKS VERSION TABLE", level=1)

    table = doc.add_table(rows=1, cols=4)
    table.style = "Table Grid"

    headers = ["Name", "Type", "Version", "Modified"]
    hdr_cells = table.rows[0].cells

    for j, h in enumerate(headers):
        hdr_cells[j].text = h
        p = hdr_cells[j].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        if p.runs:
            p.runs[0].bold = True
        shade_cell(hdr_cells[j], "D9D9D9")

    for t in tasks:
        row = table.add_row().cells
        row[0].text = t.name
        row[1].text = t.type
        row[2].text = t.version
        row[3].text = t.modified
        # opzionale: centra alcune colonne
        row[2].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        row[3].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER

    out_docx.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(out_docx))


# -------------------------
# QUICK TEST (tu cambi i path)
# -------------------------
if __name__ == "__main__":
    imgconf = r"C:\Users\claud\Desktop\streamDemo\NSPC\Imgconf.ini"
    nspc_root = r"C:\Users\claud\Desktop\streamDemo\NSPC"  # root dove risolvere i file relativi (es. FileLoaderVer)
    out = r"C:\Users\claud\Desktop\ch3_test.docx"

    tasks = read_tasks_kernel_internal(imgconf, nspc_root)
    print(tasks)
    #build_ch3_docx(tasks, out)
    #print(f"OK -> generated: {out}")
