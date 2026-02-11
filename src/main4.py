from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path
import configparser
import xml.etree.ElementTree as ET
from lxml import etree

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn



@dataclass
class TaskRow:
    name: str
    type: str
    version: str
    modified: str = "N.A."  # for now (no previous release)


#IMGCONF READER (kernel internal)
def load_imgconf_configparser(imgconf_path: Path) -> configparser.ConfigParser:

    cp = configparser.ConfigParser(
        interpolation=None,
        comment_prefixes=(";", "#", "//"),
        inline_comment_prefixes=(";", "#", "//"),
        delimiters=("=", ":"),
        strict=False,
    )
    cp.optionxform = str  # keep case

    with imgconf_path.open("r", encoding="utf-8", errors="ignore") as f:
        cp.read_file(f)

    return cp


def parse_flat_key_values(imgconf_path: Path) -> dict[str, str]:

    import re
    data: dict[str, str] = {}

    for line in imgconf_path.read_text(encoding="utf-8", errors="ignore").splitlines():
        line = line.strip()
        if not line:
            continue
        if line.startswith("//") or line.startswith("/*") or line.startswith("*"):
            continue
        # skip section headers
        if line.startswith("[") and line.endswith("]"):
            continue

        m = re.match(r"^([A-Za-z0-9_]+)\s*[:=]\s*(.+?)\s*$", line)
        if not m:
            continue
        key = m.group(1).strip()
        val = m.group(2).strip().strip('"').strip("'")
        data[key] = val

    return data


def read_tasks_kernel_internal(imgconf_path: Path) -> list[TaskRow]:
    
    if not imgconf_path.exists():
        raise FileNotFoundError(f"Imgconf.ini not found: {imgconf_path}")

    # --- Prefer section-based parsing ---
    cp = load_imgconf_configparser(imgconf_path)

    if "Settings" not in cp:
        # If your file truly has no sections, comment the cp block and use flat parsing:
        # ini = parse_flat_key_values(imgconf_path)
        raise ValueError("Missing [Settings] section in Imgconf.ini (section-based parsing).")

    s = cp["Settings"]
    tasks: list[TaskRow] = []

    # System tasks (versions in ini)
    if "FileBootVer" in s:
        tasks.append(TaskRow(name="BOOT", type="SYSTEM", version=s.get("FileBootVer", "").strip()))
    if "FileBootAPVer" in s:
        tasks.append(TaskRow(name="BOOTAP", type="SYSTEM", version=s.get("FileBootAPVer", "").strip()))
    if "FileLoaderVer" in s:
        # In legacy FileLoaderVer is a *path to a file* containing version in header.
        # For this test we just store the value itself (or you can extract from that file).
        tasks.append(TaskRow(name="Loader", type="SYSTEM", version=s.get("FileLoaderVer", "").strip()))
    if "RelKernel" in s:
        tasks.append(TaskRow(name="Kernel", type="SYSTEM", version=s.get("RelKernel", "").strip()))

    # BSP / app tasks
    n_task = int((s.get("NumTask", "0") or "0").strip())
    for i in range(1, n_task + 1):
        ttype = (s.get(f"TipoTask{i}", "") or "").strip()
        if not ttype:
            continue
        if ttype.upper() in ("NO_SCHED", "RBC"):
            continue

        version = (s.get(f"RelTask{i}", "") or "").strip()
        v1 = (s.get(f"V1_FileTask{i}", "") or "").strip()

        # Name from V1 filename stem (legacy style)
        name = Path(v1.replace("\\", "/")).stem if v1 else f"Task{i}"

        tasks.append(TaskRow(name=name, type=ttype, version=version, modified="N.A."))

    return tasks


#XML
def build_ch3_xml(tasks: list[TaskRow], out_xml: Path, title: str = "TASKS VERSION TABLE") -> None:
    doc = ET.Element("document")
    p = ET.SubElement(doc, "paragraph", number="3", title=title)
    table = ET.SubElement(p, "table", columns="Name,Type,Version,Modified")

    for t in tasks:
        ET.SubElement(
            table,
            "row",
            {
                "name": t.name or "",
                "type": t.type or "",
                "version": t.version or "",
                "modified": t.modified or "N.A.",
            },
        )

    out_xml.parent.mkdir(parents=True, exist_ok=True)
    ET.ElementTree(doc).write(out_xml, encoding="utf-8", xml_declaration=True)


# DOCX FROM XML
def _shade_cell(cell, fill_hex: str = "D9D9D9") -> None:
    tc_pr = cell._element.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill_hex)
    tc_pr.append(shd)


def _set_cell_text(cell, text: str, *, bold: bool = False, center: bool = False) -> None:
    cell.text = ""
    p = cell.paragraphs[0]
    if center:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(text if text is not None else "")
    run.bold = bold


def render_ch3_docx_from_xml(xml_path: Path, out_docx: Path) -> None:
    root = ET.parse(xml_path).getroot()
    ch3 = root.find(".//paragraph[@number='3']")
    if ch3 is None:
        raise ValueError("Chapter 3 paragraph (number='3') not found in XML")

    title = (ch3.get("title") or "TASKS VERSION TABLE").strip()
    table_node = ch3.find("./table")
    if table_node is None:
        raise ValueError("Chapter 3 <table> not found in XML")

    cols_attr = (table_node.get("columns") or "").strip()
    columns = [c.strip() for c in cols_attr.split(",") if c.strip()] or ["Name", "Type", "Version", "Modified"]

    doc = Document()
    doc.add_heading(f"3. {title}", level=1)
    doc.add_paragraph("")

    table = doc.add_table(rows=1, cols=len(columns))
    table.style = "Table Grid"

    # Header
    hdr = table.rows[0].cells
    for i, col in enumerate(columns):
        _set_cell_text(hdr[i], col, bold=True, center=True)
        _shade_cell(hdr[i], "D9D9D9")

    # Rows
    for row in table_node.findall("./row"):
        name = (row.get("name") or "").strip()
        typ = (row.get("type") or "").strip()
        ver = (row.get("version") or "").strip()
        mod = (row.get("modified") or "N.A.").strip()

        r = table.add_row().cells
        _set_cell_text(r[0], name, center=False)
        _set_cell_text(r[1], typ, center=True)
        _set_cell_text(r[2], ver, center=True)
        _set_cell_text(r[3], mod, center=True)

    out_docx.parent.mkdir(parents=True, exist_ok=True)
    doc.save(out_docx)


 # Adjust these paths:
IMGCONF = Path(r"C:\Users\cpisa\Desktop\stream\DEVRASTA\NSPC\Imgconf.ini")

OUT_XML = Path(r"C:\Users\cpisa\Desktop\out\ch3_test.xml")
OUT_DOCX = Path(r"C:\Users\cpisa\Desktop\out\ch3_test.docx")

# 1) Read
tasks = read_tasks_kernel_internal(IMGCONF)
print(f"Loaded {len(tasks)} tasks from Imgconf.ini")

# 2) XML
build_ch3_xml(tasks, OUT_XML, title="TASK VERSION IDENTIFICATION")
print(f"XML written: {OUT_XML}")

# 3) DOCX from XML
render_ch3_docx_from_xml(OUT_XML, OUT_DOCX)
print(f"DOCX written: {OUT_DOCX}")